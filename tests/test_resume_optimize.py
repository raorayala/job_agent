"""Unit tests for AI Optimize resume suggestions and profile optimize."""

from pathlib import Path

from docx import Document

from job_agent.config import Settings
from job_agent.models import CandidateProfile, MatchExplanation, ParsedJob, Recommendation
from job_agent.profile_optimize import analyze_profile_against_target, apply_skills_order_to_profile
from job_agent.resume_optimize import (
    apply_accepted_suggestions_to_draft,
    compute_keyword_gap,
    extract_jd_keywords,
    filter_ats_keywords,
    generate_optimize_suggestions,
    is_actionable_ats_keyword,
    load_optimize_bundle,
    sanitize_job_description_for_keywords,
    update_suggestion,
)


def _settings(tmp_path: Path) -> Settings:
    return Settings(
        project_root=tmp_path,
        master_resume_path=tmp_path / "master.docx",
        jobs_applied_folder=tmp_path / "Jobs Applied",
        jobs_draft_folder=tmp_path / "_drafts",
        preferred_browser="system",
        gmail_credentials_path=tmp_path / "creds.json",
        gmail_token_path=tmp_path / "token.json",
        gmail_search_query="",
        gmail_label=None,
        imap_host="outlook.office365.com",
        imap_port=993,
        imap_username="",
        imap_password="",
        imap_folder="INBOX",
        min_match_score=60.0,
        llm_provider="none",
        ollama_base_url="",
        ollama_model="",
        database_path=tmp_path / "test.db",
        log_level="INFO",
        config_path=tmp_path / "config.yaml",
    )


def test_sanitize_strips_urls_and_email_boilerplate():
    raw = (
        "Murali, don't miss these jobs. "
        "Python Docker required. "
        "https://cts.indeed.com/v3/h4siaaaaaaaa?utm=1 "
        "Click here to apply now."
    )
    cleaned = sanitize_job_description_for_keywords(raw)
    assert "cts.indeed.com" not in cleaned.lower()
    assert "https://" not in cleaned.lower()
    assert "don't miss" not in cleaned.lower()
    assert "Click here to apply" not in cleaned
    assert "Python" in cleaned
    assert "Docker" in cleaned


def test_extract_jd_keywords_ignores_noise_and_benefits():
    jd = (
        "Senior Engineer III / Sr level 3. "
        "Need Python, Docker, SQL, AWS, Microservices, and CI/CD. "
        "Benefits include insurance, leave, 401k match, paid PTO, health, medical, vision, dental. "
        "Equal opportunity employer. Background check required. "
        "Apply via https://cts.indeed.com/v3/h4siaaaaaaaa tracking link."
    )
    terms = [t.lower() for t, _ in extract_jd_keywords(jd, limit=30)]
    for noise in (
        "insurance",
        "leave",
        "match",
        "paid",
        "pto",
        "health",
        "medical",
        "vision",
        "dental",
        "background",
        "equal",
        "opportunity",
        "iii",
        "sr",
        "indeed",
        "cts",
        "h4siaaaaaaaa",
    ):
        assert noise not in terms, f"unexpected noise keyword: {noise}"
    for skill in ("python", "docker", "sql", "aws", "microservices"):
        assert skill in terms, f"missing technical keyword: {skill}"
    assert any(t in {"ci/cd", "ci", "cd"} for t in terms) or "ci/cd" in " ".join(terms)


def test_is_actionable_ats_keyword_rules():
    assert is_actionable_ats_keyword("Python")
    assert is_actionable_ats_keyword("C++")
    assert is_actionable_ats_keyword("SRE")
    assert is_actionable_ats_keyword("CI/CD")
    assert not is_actionable_ats_keyword("iii")
    assert not is_actionable_ats_keyword("insurance")
    assert not is_actionable_ats_keyword("leave")
    assert not is_actionable_ats_keyword("sr")
    assert not is_actionable_ats_keyword("h4siaaaaaaaa")
    assert not is_actionable_ats_keyword("cts.indeed.com")


def test_compute_keyword_gap_filters_noise_from_suggestions_path():
    job = ParsedJob(
        title="Backend Engineer",
        company="Acme",
        description=(
            "Python Kubernetes AWS required. Benefits: insurance leave match paid PTO. "
            "cts.indeed.com/v3/h4siaaaaaaaa iii sr"
        ),
    )
    match = MatchExplanation(
        score=70.0,
        recommendation=Recommendation.WORTH_REVIEWING,
        matched_skills=["Python"],
        missing_skills=["Kubernetes", "insurance", "leave", "iii"],
        summary="partial",
    )
    matched, gaps, readiness = compute_keyword_gap(
        job,
        resume_text="Python developer with APIs",
        match=match,
    )
    lowered_gaps = {g.lower() for g in gaps}
    lowered_matched = {m.lower() for m in matched}
    assert "python" in lowered_matched
    assert "kubernetes" in lowered_gaps or "kubernetes" in lowered_matched
    assert "insurance" not in lowered_gaps
    assert "leave" not in lowered_gaps
    assert "iii" not in lowered_gaps
    assert 0 <= readiness <= 100
    assert filter_ats_keywords(["Python", "insurance", "iii", "Docker"]) == ["Python", "Docker"]


def test_generate_optimize_suggestions_and_accept_apply(tmp_path: Path):
    master = tmp_path / "master.docx"
    doc = Document()
    doc.add_heading("Jane Doe", level=1)
    doc.add_paragraph("Senior Python Developer with Docker and AWS experience.")
    doc.add_paragraph("- Built APIs with Python and FastAPI")
    doc.add_paragraph("- Deployed services using Docker")
    doc.save(str(master))

    draft = tmp_path / "_drafts" / "Acme_Engineer_Draft_v1.docx"
    draft.parent.mkdir(parents=True, exist_ok=True)
    Document(str(master)).save(str(draft))

    job = ParsedJob(
        title="Senior Engineer",
        company="Acme",
        description="Looking for Python Docker Kubernetes AWS FastAPI engineer with strong leadership.",
    )
    match = MatchExplanation(
        score=78.0,
        recommendation=Recommendation.WORTH_REVIEWING,
        matched_skills=["Python", "Docker"],
        missing_skills=["Kubernetes"],
        summary="Partial match",
    )
    settings = _settings(tmp_path)
    settings.master_resume_path = master

    bundle = generate_optimize_suggestions(
        job=job,
        match=match,
        master_resume_path=master,
        draft_resume_path=draft,
        job_id=1,
        settings=settings,
    )
    assert bundle.suggestions
    assert bundle.readiness_score >= 0
    assert load_optimize_bundle(draft) is not None
    noise = {"insurance", "leave", "iii", "pto", "indeed"}
    assert not noise.intersection({g.lower() for g in bundle.keyword_gaps})
    for suggestion in bundle.suggestions:
        if suggestion.kind == "keyword_insert" and suggestion.keyword:
            assert is_actionable_ats_keyword(suggestion.keyword)

    # Accept first non-confirm suggestion if possible
    target = next(
        (s for s in bundle.suggestions if not s.suggested_text.startswith("[CONFIRM BEFORE ADDING]")),
        bundle.suggestions[0],
    )
    updated = update_suggestion(draft, target.id, "accept")
    assert any(s.id == target.id and s.status == "accepted" for s in updated.suggestions)

    applied = apply_accepted_suggestions_to_draft(draft, master_resume_path=master)
    assert applied.exists()
    text = "\n".join(p.text for p in Document(str(applied)).paragraphs)
    assert "AI Optimize" in text or target.suggested_text[:20] in text


def test_profile_optimize_keyword_gap_and_skills_audit():
    profile = CandidateProfile(
        target_titles=["Staff Backend Engineer"],
        required_skills=["Python", "Docker", "PostgreSQL"],
        preferred_skills=["Kafka"],
        years_experience=8,
    )
    result = analyze_profile_against_target(
        profile,
        job_description="Staff Backend Engineer Python Docker Kubernetes AWS PostgreSQL Kafka leadership",
        resume_text="Python Docker PostgreSQL experience building APIs",
    )
    assert 0 <= result.readiness_score <= 100
    assert result.headline
    assert len(result.headline) <= 220
    assert result.about_summary
    assert result.skills_audit
    assert result.suggested_skills_order

    updated = apply_skills_order_to_profile(profile, result.suggested_skills_order, accepted_adds=["Kubernetes"])
    assert "Python" in updated.required_skills

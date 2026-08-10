"""Unit tests for top 10 platforms, review-first resume approval workflow, and browser behavior."""

import json
from pathlib import Path
import pytest
from sqlalchemy import select

from job_agent.application_tracker import (
    approve_resume_draft_for_job,
    create_resume_draft_for_job,
    record_parsed_job,
)
from job_agent.config import Settings, validate_config
from job_agent.database import JobRecord, init_db
from job_agent.models import CandidateProfile, ParsedJob
from job_agent.platform_fetcher import (
    TOP_10_PLATFORMS,
    generate_platform_search_urls,
    identify_platform_from_url,
    import_job_from_url,
)
from job_agent.cli import _open_in_browser


def test_top_10_platforms_urls_and_identification():
    urls = generate_platform_search_urls("Software Engineer", "Remote")
    assert len(urls) == 10
    for p in TOP_10_PLATFORMS:
        assert p in urls
        assert "14" in urls[p] or "r1209600" in urls[p] or "jobs" in urls[p]

    assert identify_platform_from_url("https://www.linkedin.com/jobs/view/123") == "linkedin"
    assert identify_platform_from_url("https://www.indeed.com/viewjob?jk=abc") == "indeed"
    assert identify_platform_from_url("https://www.glassdoor.com/job-listing/x") == "glassdoor"
    assert identify_platform_from_url("https://wellfound.com/jobs/123") == "wellfound"


def test_import_job_from_url_auto_parser():
    parsed = import_job_from_url("https://www.linkedin.com/jobs/view/999888")
    assert isinstance(parsed, ParsedJob)
    assert parsed.job_url == "https://www.linkedin.com/jobs/view/999888"
    assert parsed.source_platform == "linkedin"
    assert parsed.title != ""


def test_config_validation():
    warnings = validate_config({"profile": {"target_titles": [], "required_skills": []}})
    assert len(warnings) >= 2
    assert any("target_titles" in w for w in warnings)
    assert any("required_skills" in w for w in warnings)


def test_review_first_resume_approval_workflow(tmp_path):
    db_path = tmp_path / "test_workflow.db"
    drafts_dir = tmp_path / "_drafts"
    applied_dir = tmp_path / "Jobs Applied"
    master_docx = tmp_path / "master.docx"

    # Create dummy master resume docx
    from docx import Document
    doc = Document()
    doc.add_heading("John Doe", level=1)
    doc.add_paragraph("Senior Python Developer with 8 years experience.")
    doc.save(str(master_docx))

    SessionLocal = init_db(db_path)
    session = SessionLocal()

    profile = CandidateProfile(
        target_titles=["Senior Engineer"],
        required_skills=["Python", "Docker"],
        master_resume_path=str(master_docx),
    )
    settings = Settings(
        project_root=tmp_path,
        master_resume_path=master_docx,
        jobs_applied_folder=applied_dir,
        jobs_draft_folder=drafts_dir,
        preferred_browser="system",
        gmail_credentials_path=tmp_path / "creds.json",
        gmail_token_path=tmp_path / "token.json",
        gmail_search_query="",
        gmail_label=None,
        min_match_score=60.0,
        llm_provider="none",
        ollama_base_url="",
        ollama_model="",
        database_path=db_path,
        log_level="INFO",
        config_path=tmp_path / "config.yaml",
    )

    try:
        # 1. Record job
        parsed_job = ParsedJob(
            title="Senior Engineer",
            company="Acme Testing Corp",
            job_url="https://example.com/job/101",
            description="Looking for Python, Docker, Kubernetes engineer.",
            source_platform="indeed",
        )
        record, match, _ = record_parsed_job(session, parsed_job, profile)
        assert record.id is not None

        # 2. Generate Draft Resume
        job, draft_p, summary_p, diff_text = create_resume_draft_for_job(
            session=session,
            job_id=record.id,
            profile=profile,
            settings=settings,
        )

        assert draft_p.exists()
        assert "_drafts" in str(draft_p)
        assert job.status == "Resume draft ready"
        assert job.approval_status == "awaiting_review"

        # SAFEGUARD VERIFICATION: Final applied resume folder MUST NOT exist yet before explicit approval!
        from job_agent.document_exporter import resume_filename
        out_filename = resume_filename("Acme Testing Corp", "Senior Engineer")
        final_expected_path = applied_dir / "Acme_Testing_Corp" / "Senior_Engineer" / out_filename
        assert not final_expected_path.exists()

        # 3. Explicit User Approval
        approved_job, final_path = approve_resume_draft_for_job(
            session=session,
            job_id=record.id,
            settings=settings,
        )

        assert approved_job.status == "Approved"
        assert approved_job.approval_status == "approved"
        assert final_path.exists()
        assert str(final_path) == str(final_expected_path)

    finally:
        session.close()


def test_browser_launcher_default_behavior(monkeypatch):
    called_urls = []

    def mock_open_new_tab(url):
        called_urls.append(url)
        return True

    import webbrowser
    import job_agent.cli as cli_mod

    # Force fallback path (no Chrome binary) to verify webbrowser.new_tab behavior
    monkeypatch.setattr(cli_mod, "_find_chrome_path", lambda: None)
    monkeypatch.setattr(webbrowser, "open_new_tab", mock_open_new_tab)

    _open_in_browser("https://www.linkedin.com/jobs/search", browser_choice="system", new_tab=True)
    assert len(called_urls) == 1
    assert called_urls[0] == "https://www.linkedin.com/jobs/search"


def test_browser_launcher_opens_chrome_new_tab(monkeypatch, tmp_path):
    launched = []
    fake_chrome = tmp_path / "chrome.exe"
    fake_chrome.write_text("", encoding="utf-8")

    import job_agent.cli as cli_mod
    import subprocess

    monkeypatch.setattr(cli_mod, "_find_chrome_path", lambda: str(fake_chrome))

    def mock_popen(cmd, *args, **kwargs):
        launched.append(cmd)
        class Dummy:
            pass
        return Dummy()

    monkeypatch.setattr(subprocess, "Popen", mock_popen)

    _open_in_browser("http://localhost:8000/", browser_choice="chrome", new_tab=True)
    assert len(launched) == 1
    assert launched[0][0] == str(fake_chrome)
    assert "--new-tab" in launched[0]
    assert "http://localhost:8000/" in launched[0]

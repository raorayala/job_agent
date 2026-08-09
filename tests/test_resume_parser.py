"""Tests for resume text extraction from DOCX files."""

from __future__ import annotations

from pathlib import Path
from docx import Document

from job_agent.resume_parser import extract_resume_text


def test_extract_resume_text_success(tmp_path: Path) -> None:
    doc_path = tmp_path / "sample_resume.docx"
    doc = Document()
    doc.add_heading("Jane Developer", level=1)
    doc.add_paragraph("Experienced Senior Python Developer with expertise in SQL, Docker, and AWS.")

    table = doc.add_table(rows=1, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Skills"
    hdr_cells[1].text = "Python, Java, SpringBoot"

    doc.save(str(doc_path))

    extracted = extract_resume_text(doc_path)
    assert "Jane Developer" in extracted
    assert "Experienced Senior Python Developer" in extracted
    assert "Python, Java, SpringBoot" in extracted


def test_extract_resume_text_nonexistent() -> None:
    text = extract_resume_text("C:/nonexistent/path/resume.docx")
    assert text == ""

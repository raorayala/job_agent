"""Tests for resume tailoring and DOCX export."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from job_agent.models import MatchExplanation, ParsedJob, Recommendation
from job_agent.resume_tailor import tailor_resume


def test_tailor_resume_creates_docx(tmp_path: Path) -> None:
    # Create a dummy master resume DOCX
    master_docx = tmp_path / "master_resume.docx"
    doc = Document()
    doc.add_heading("Jane Doe - Software Engineer", level=1)
    doc.add_paragraph("Experienced Python Developer with 5 years building web applications.")
    doc.save(str(master_docx))

    output_dir = tmp_path / "Desktop" / "Jobs Applied"

    job = ParsedJob(
        title="Backend Python Developer",
        company="Acme Corp",
        source_platform="indeed",
        job_url="https://indeed.com/viewjob?jk=777",
        description="Looking for Python, SQL, Docker",
    )

    match = MatchExplanation(
        score=85.0,
        recommendation=Recommendation.STRONG_MATCH,
        matched_skills=["Python", "SQL"],
        missing_skills=["Docker"],
        concerns=[],
        summary="Score 85/100 (Strong match).",
    )

    res_path = tailor_resume(
        job=job,
        match=match,
        master_resume_path=master_docx,
        output_base_dir=output_dir,
        dry_run=False,
    )

    assert res_path.exists()
    assert res_path.suffix == ".docx"
    assert "Acme_Corp" in res_path.name or "Acme" in res_path.name

    # Verify sidecar summary text file was written
    summary_txt = res_path.parent / "Application_Summary.txt"
    assert summary_txt.exists()
    content = summary_txt.read_text(encoding="utf-8")
    assert "Backend Python Developer" in content
    assert "Acme Corp" in content

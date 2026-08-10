"""Tests for review-gated Auto Apply and LinkedIn Optimization."""

from pathlib import Path

from docx import Document

from job_agent.auto_apply_service import evaluate_auto_apply_eligibility, launch_auto_apply
from job_agent.config import Settings
from job_agent.database import JobRecord, init_db
from job_agent.linkedin_optimize import optimize_linkedin_profile
from job_agent.models import ApplicationStatus
from job_agent.resume_optimize import OptimizeBundle, OptimizeSuggestion, save_optimize_bundle


def _settings(tmp_path: Path) -> Settings:
    return Settings(
        project_root=tmp_path,
        master_resume_path=tmp_path / "master.docx",
        jobs_applied_folder=tmp_path / "Jobs Applied",
        jobs_draft_folder=tmp_path / "_drafts",
        preferred_browser="system",
        gmail_credentials_path=tmp_path / "creds.json",
        gmail_token_path=tmp_path / "token.json",
        gmail_search_query="",
        gmail_label=None,
        min_match_score=60.0,
        llm_provider="none",
        ollama_base_url="",
        ollama_model="",
        database_path=tmp_path / "test.db",
        log_level="INFO",
        config_path=tmp_path / "config.yaml",
    )


def test_auto_apply_blocked_until_approved(tmp_path: Path):
    settings = _settings(tmp_path)
    SessionLocal = init_db(settings.database_path)
    session = SessionLocal()
    job = JobRecord(
        title="Engineer",
        company="Acme",
        job_url="https://example.com/jobs/1",
        status=ApplicationStatus.SAVED.value,
        approval_status="awaiting_review",
    )
    session.add(job)
    session.commit()
    session.refresh(job)

    eligibility = evaluate_auto_apply_eligibility(job)
    assert eligibility.eligible is False
    assert any("Approve" in b for b in eligibility.blockers)


def test_auto_apply_launch_after_approval(tmp_path: Path, monkeypatch):
    settings = _settings(tmp_path)
    final = tmp_path / "Jobs Applied" / "Acme" / "Engineer" / "resume.docx"
    final.parent.mkdir(parents=True, exist_ok=True)
    Document().save(str(final))
    draft = tmp_path / "_drafts" / "draft.docx"
    draft.parent.mkdir(parents=True, exist_ok=True)
    Document().save(str(draft))

    bundle = OptimizeBundle(
        job_id=1,
        draft_resume_path=str(draft),
        suggestions=[
            OptimizeSuggestion(
                id="1",
                kind="keyword_insert",
                status="accepted",
                original_text="Built APIs",
                suggested_text="Built APIs with Python",
                rationale="ATS keyword",
            )
        ],
    )
    save_optimize_bundle(bundle, draft)

    SessionLocal = init_db(settings.database_path)
    session = SessionLocal()
    job = JobRecord(
        title="Engineer",
        company="Acme",
        job_url="https://example.com/jobs/1",
        status=ApplicationStatus.APPROVED.value,
        approval_status="approved",
        final_resume_path=str(final),
        draft_resume_path=str(draft),
        diff_summary="AI Optimize suggestions applied",
    )
    session.add(job)
    session.commit()
    session.refresh(job)

    monkeypatch.setattr("job_agent.auto_apply_service.open_url_in_browser", lambda *a, **k: True)
    monkeypatch.setattr("job_agent.auto_apply_service.open_path_in_explorer", lambda *a, **k: True)

    eligibility = evaluate_auto_apply_eligibility(job)
    assert eligibility.eligible is True
    assert eligibility.ats_ready is True

    result = launch_auto_apply(
        session,
        job.id,
        settings,
        open_browser=True,
        open_resume_folder=True,
        mark_as_applied=False,
        confirm=False,
    )
    assert result.status == "success"
    assert result.marked_status == ApplicationStatus.READY_TO_APPLY.value
    session.refresh(job)
    assert job.status == ApplicationStatus.READY_TO_APPLY.value


def test_linkedin_optimize_returns_headline(tmp_path: Path, monkeypatch):
    master = tmp_path / "master.docx"
    Document().save(str(master))
    # Write minimal config.yaml for profile loader
    cfg = tmp_path / "config.yaml"
    cfg.write_text(
        "profile:\n  target_titles: [Staff Engineer]\n  required_skills: [Python, Docker]\n  preferred_skills: [Kafka]\n  years_experience: 8\n",
        encoding="utf-8",
    )
    monkeypatch.setenv("CONFIG_PATH", str(cfg))
    monkeypatch.setenv("MASTER_RESUME_PATH", str(master))
    monkeypatch.setenv("LLM_PROVIDER", "none")

    result = optimize_linkedin_profile(target_role="Staff Backend Engineer", job_description="Python Docker Kafka AWS")
    assert result["status"] == "success"
    assert result["headline"]
    assert len(result["headline"]) <= 220
    assert result["about_summary"]
    assert "linkedin_tips" in result

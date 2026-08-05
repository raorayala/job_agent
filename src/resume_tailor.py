"""Tailor resume DOCX for ATS keyword alignment."""

from __future__ import annotations

import re
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from docx import Document

from src.config import get_env, get_jobs_applied_folder
from src.job_analyzer import MatchResult
from src.job_parser import ParsedJob


def _safe_filename(text: str) -> str:
    cleaned = re.sub(r"[^\w\s-]", "", text).strip().replace(" ", "_")
    return cleaned[:80] or "job"


def tailor_resume(
    job: ParsedJob,
    match: MatchResult,
    master_resume_path: str | None = None,
) -> Path:
    """
    Create a tailored resume copy with an ATS keyword summary section.

    Phase 1: appends a keyword block. Later phases can reorder bullets via LLM.
    """
    master_path = Path(master_resume_path or get_env("MASTER_RESUME_PATH"))
    if not master_path.exists():
        raise FileNotFoundError(
            f"Master resume not found at {master_path}. "
            "Place your master resume there or set MASTER_RESUME_PATH in .env"
        )

    doc = Document(str(master_path))

    # Append ATS alignment section (easy to remove manually before applying)
    doc.add_page_break()
    doc.add_heading("ATS Keyword Alignment (auto-generated — review before submitting)", level=2)
    doc.add_paragraph(f"Target role: {job.title} at {job.company}")
    doc.add_paragraph(f"Match score: {match.score:.0f}/100")
    doc.add_paragraph(f"Matched skills: {', '.join(match.matched_skills) or 'None detected'}")
    doc.add_paragraph(
        "Suggested keywords to weave into your summary and experience bullets: "
        + ", ".join(sorted(set(job.required_skills + job.preferred_skills + match.matched_skills)))
    )

    output_dir = get_jobs_applied_folder() / "tailored_resumes"
    timestamp = datetime.now().strftime("%Y%m%d")
    filename = f"{timestamp}_{_safe_filename(job.company)}_{_safe_filename(job.title)}.docx"
    output_path = output_dir / filename

    doc.save(str(output_path))
    return output_path


def export_application_summary(job: ParsedJob, match: MatchResult, resume_path: Path) -> Path:
    """Write a plain-text summary alongside the tailored resume."""
    exports = get_jobs_applied_folder() / "exports"
    summary_path = exports / (resume_path.stem + ".txt")
    summary_path.write_text(
        "\n".join(
            [
                f"Title: {job.title}",
                f"Company: {job.company}",
                f"Platform: {job.platform}",
                f"URL: {job.source_url}",
                f"Match score: {match.score:.0f}",
                f"Resume: {resume_path}",
                f"Analysis: {match.summary}",
            ]
        ),
        encoding="utf-8",
    )
    return summary_path

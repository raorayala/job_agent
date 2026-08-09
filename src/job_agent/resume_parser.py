"""Utility for extracting text from master resume DOCX files."""

from __future__ import annotations

from pathlib import Path
from docx import Document

from job_agent.logging_config import get_logger

logger = get_logger(__name__)


def extract_resume_text(docx_path: Path | str | None) -> str:
    """
    Extract all readable plain text from a DOCX resume file including paragraphs and tables.
    Returns empty string if file does not exist or cannot be parsed.
    """
    if not docx_path:
        return ""

    path = Path(docx_path)
    if not path.exists():
        logger.warning("Master resume file not found at: %s", path)
        return ""

    try:
        doc = Document(str(path))
        text_parts: list[str] = []

        # Extract text from paragraphs
        for p in doc.paragraphs:
            if p.text and p.text.strip():
                text_parts.append(p.text.strip())

        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                if row_text:
                    text_parts.append(" | ".join(row_text))

        full_text = "\n".join(text_parts)
        logger.info("Extracted %d characters from master resume: %s", len(full_text), path)
        return full_text
    except Exception as exc:
        logger.error("Error extracting text from master resume DOCX %s: %s", path, exc)
        return ""

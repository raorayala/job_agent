"""Truthful AI Optimize: keyword/bullet/summary suggestions with accept/reject/edit."""

from __future__ import annotations

import json
import re
import uuid
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document

from job_agent.config import Settings, get_settings
from job_agent.job_normalizer import normalize_text
from job_agent.logging_config import get_logger
from job_agent.models import MatchExplanation, ParsedJob
from job_agent.ollama_service import generate_ollama_completion
from job_agent.resume_parser import extract_resume_text

logger = get_logger(__name__)

_BULLET_RE = re.compile(r"^[\u2022\-\*\u25CF\u25E6]\s*(.+)$")
_WORD_RE = re.compile(r"[A-Za-z][A-Za-z0-9+.#/-]{1,}")
_URL_RE = re.compile(
    r"(?i)\b(?:https?://|www\.)\S+|"
    r"\b[\w.-]+\.(?:com|net|org|io|co|edu|gov|info|biz|us)(?:/\S*)?"
)
_EMAIL_ADDR_RE = re.compile(r"(?i)\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b")
_BOILERPLATE_RE = re.compile(
    r"(?i)\b(?:"
    r"don'?t miss these jobs?|"
    r"click here to apply|"
    r"view(?:\s+this)?\s+job|"
    r"apply(?:\s+now)?|"
    r"unsubscribe|"
    r"view in browser|"
    r"you(?:'re| are) receiving this|"
    r"manage (?:email )?preferences|"
    r"privacy policy|"
    r"equal opportunity employer|"
    r"affirmative action|"
    r"all qualified applicants"
    r")\b[^.!?\n]*[.!?]?"
)
_REPEATED_CHAR_RE = re.compile(r"(.)\1{3,}")
_TRACKING_TOKEN_RE = re.compile(r"(?i)^(?=[a-z0-9]*\d)[a-z0-9]{10,}$")
_LEVEL_TOKEN_RE = re.compile(r"(?i)^(?:level|lvl)_?\d*$")

# Short/tech tokens that should survive length and pattern filters.
TECH_TOKEN_ALLOWLIST = {
    "c", "c++", "c#", "go", "r", "js", "ts", "py", "sql", "aws", "gcp", "azure",
    "ci", "cd", "ci/cd", "sre", "ml", "ai", "ui", "ux", "api", "sdk", "ide",
    "k8s", "eks", "ecs", "etl", "elt", "orm", "rest", "grpc", "ssl",
    "tls", "vpn", "dns", "tcp", "udp", "db", "nosql", "nlp", "llm", "rag",
    "qa", "sdet", "git", "svn", "npm", "pip", "jvm", "jdk", "jpa", "jdbc",
    "css", "html", "xml", "json", "yaml", "toml", "bash", "zsh", "ios", "android",
    "dotnet", ".net", "node", "react", "vue", "angular", "kafka", "redis", "mongo",
    "postgres", "mysql", "oracle", "snowflake", "databricks", "spark", "hadoop",
    "terraform", "ansible", "docker", "kubernetes", "linux", "unix", "python",
    "java", "scala", "kotlin", "rust", "ruby", "php", "swift", "typescript",
    "javascript", "golang", "fastapi", "django", "flask", "spring", "rails",
}

_ROMAN_NUMERALS = {
    "i", "ii", "iii", "iv", "v", "vi", "vii", "viii", "ix", "x", "xi", "xii",
    "xiii", "xiv", "xv",
}

# Generic English + HR/benefits + job-board/noise terms excluded from ATS gaps.
ATS_STOP_WORDS = {
    # English function words
    "and", "the", "for", "with", "you", "our", "will", "are", "this", "that", "from",
    "have", "your", "able", "work", "team", "role", "job", "jobs", "years", "experience",
    "including", "using", "required", "preferred", "strong", "must", "should",
    "about", "into", "such", "across", "within", "other", "more", "than", "over",
    "into", "their", "they", "them", "who", "what", "when", "where", "which", "while",
    "also", "any", "all", "can", "may", "not", "but", "per", "via", "etc", "new",
    "well", "based", "related", "looking", "join", "please", "click", "here", "apply",
    "view", "miss", "these", "those", "make", "sure", "best", "great", "good",
    # Job level / roman / title noise
    "sr", "jr", "snr", "jnr", "mid", "level", "lvl", "senior", "junior", "staff",
    "principal", "lead", "intern", "internship", "fulltime", "parttime", "contract",
    # HR / benefits / legal boilerplate
    "insurance", "leave", "match", "401k", "401", "paid", "pto", "health", "medical",
    "vision", "dental", "background", "equal", "opportunity", "employer", "benefits",
    "vacation", "sick", "bonus", "salary", "compensation", "retirement", "disability",
    "parental", "maternity", "paternity", "holiday", "holidays", "perks", "wellness",
    "stipend", "reimbursement", "eoe", "eeoc", "affirmative", "action", "diversity",
    "inclusion", "equity", "belonging", "accommodation", "disability", "veteran",
    "status", "race", "gender", "religion", "orientation", "identity", "protected",
    "qualified", "applicants", "receive", "consideration", "without", "regard",
    "flexible", "remote", "hybrid", "onsite", "relocation", "sponsorship", "visa",
    "unlimited", "time", "off", "package", "competitive", "commensurate",
    # Job boards / tracking / contact chrome
    "indeed", "glassdoor", "ziprecruiter", "linkedin", "dice", "monster", "lever",
    "greenhouse", "workday", "taleo", "icims", "smartrecruiters", "com", "www",
    "http", "https", "mailto", "email", "phone", "contact", "unsubscribe", "browser",
    "preferences", "privacy", "policy", "cts", "utm", "tracking",
}


def sanitize_job_description_for_keywords(text: str) -> str:
    """Strip URLs, emails, and recruiting/email boilerplate before keyword mining."""
    cleaned = text or ""
    cleaned = _URL_RE.sub(" ", cleaned)
    cleaned = _EMAIL_ADDR_RE.sub(" ", cleaned)
    cleaned = _BOILERPLATE_RE.sub(" ", cleaned)
    # Drop common personalized openers like "Name, don't miss..."
    cleaned = re.sub(r"(?im)^[A-Z][A-Za-z'’.-]{1,30},\s+", " ", cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned)
    return cleaned.strip()


def is_actionable_ats_keyword(token: str) -> bool:
    """True when a token is a plausible technical skill / competency for gap analysis."""
    raw = (token or "").strip()
    if not raw:
        return False
    key = raw.lower()

    if key in TECH_TOKEN_ALLOWLIST or raw in TECH_TOKEN_ALLOWLIST:
        return True
    if any(ch in raw for ch in "+#/") and len(key) >= 2 and key not in ATS_STOP_WORDS:
        return True

    if key in ATS_STOP_WORDS:
        return False
    if key in _ROMAN_NUMERALS or re.fullmatch(r"[ivxlcdm]+", key):
        return False
    if len(key) < 3:
        return False
    if _LEVEL_TOKEN_RE.fullmatch(key) or key in {"sr", "jr", "snr", "jnr"}:
        return False
    if re.search(r"\.(?:com|net|org|io|co|edu|gov)\b", key) or key.endswith(
        (".com", ".net", ".org", ".io")
    ):
        return False
    if _REPEATED_CHAR_RE.search(key):
        return False
    if _TRACKING_TOKEN_RE.fullmatch(key):
        return False
    # Reject bare version-ish tokens like v3 / l2 unless allowlisted.
    if re.fullmatch(r"[a-z]\d{1,3}", key):
        return False
    return True


def filter_ats_keywords(terms: list[str] | tuple[str, ...]) -> list[str]:
    """Dedupe and drop non-actionable ATS terms while preserving order."""
    seen: set[str] = set()
    out: list[str] = []
    for term in terms:
        clean = (term or "").strip()
        if not clean:
            continue
        key = clean.lower()
        if key in seen or not is_actionable_ats_keyword(clean):
            continue
        seen.add(key)
        out.append(clean)
    return out


def _skill_in_text(skill: str, text: str) -> bool:
    clean = skill.strip()
    if not clean or not text:
        return False
    pattern = re.compile(r"\b" + re.escape(clean) + r"\b", re.IGNORECASE)
    return bool(pattern.search(text))


def extract_jd_keywords(job_description: str, *, limit: int = 25) -> list[tuple[str, int]]:
    """Rank high-frequency technical-looking tokens from a job description."""
    cleaned = sanitize_job_description_for_keywords(job_description or "")
    counts: dict[str, int] = {}
    originals: dict[str, str] = {}
    for token in _WORD_RE.findall(cleaned):
        if not is_actionable_ats_keyword(token):
            continue
        key = token.lower()
        counts[key] = counts.get(key, 0) + 1
        # Prefer display form with technical punctuation / common casing.
        prev = originals.get(key)
        if prev is None or any(ch in token for ch in "+#./") or (
            token[0].isupper() and not prev[0].isupper()
        ):
            originals[key] = token
    ranked = sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))
    return [(originals.get(term, term), freq) for term, freq in ranked[:limit]]


def compute_keyword_gap(
    job: ParsedJob,
    resume_text: str,
    match: MatchExplanation | None = None,
) -> tuple[list[str], list[str], float]:
    """Return (matched_keywords, gap_keywords, readiness_score 0-100)."""
    jd = f"{job.title} {job.company} {job.description}"
    resume_norm = normalize_text(resume_text)
    ranked = extract_jd_keywords(jd, limit=30)
    matched: list[str] = []
    gaps: list[str] = []

    for term, _freq in ranked:
        if _skill_in_text(term, resume_norm) or term.lower() in resume_norm:
            matched.append(term)
        else:
            gaps.append(term)

    if match:
        for skill in filter_ats_keywords(match.matched_skills):
            if skill.lower() not in {m.lower() for m in matched}:
                matched.append(skill)
        for skill in filter_ats_keywords(match.missing_skills):
            # Missing from profile overlap; still a gap if absent from resume
            if not _skill_in_text(skill, resume_norm):
                if skill.lower() not in {g.lower() for g in gaps}:
                    gaps.append(skill)

    matched = filter_ats_keywords(matched)
    gaps = filter_ats_keywords(gaps)

    total = max(len(matched) + len(gaps), 1)
    readiness = round(100.0 * len(matched) / total, 1)
    if match is not None:
        readiness = round((readiness * 0.4) + (float(match.score) * 0.6), 1)
    return matched[:20], gaps[:20], readiness


@dataclass
class OptimizeSuggestion:
    """One editable optimize suggestion for a resume draft."""

    id: str
    kind: str  # keyword_insert | bullet_rewrite | summary_rewrite
    status: str  # pending | accepted | rejected | edited
    original_text: str
    suggested_text: str
    rationale: str
    target_section: str = ""
    keyword: str = ""

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> OptimizeSuggestion:
        return cls(
            id=str(data.get("id") or uuid.uuid4()),
            kind=str(data.get("kind") or "bullet_rewrite"),
            status=str(data.get("status") or "pending"),
            original_text=str(data.get("original_text") or ""),
            suggested_text=str(data.get("suggested_text") or ""),
            rationale=str(data.get("rationale") or ""),
            target_section=str(data.get("target_section") or ""),
            keyword=str(data.get("keyword") or ""),
        )


@dataclass
class OptimizeBundle:
    """Persisted optimize session for one draft."""

    job_id: int | None = None
    draft_resume_path: str = ""
    created_at: str = ""
    source: str = "rule"  # rule | ollama | hybrid
    readiness_score: float = 0.0
    keyword_gaps: list[str] = field(default_factory=list)
    matched_keywords: list[str] = field(default_factory=list)
    suggestions: list[OptimizeSuggestion] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "job_id": self.job_id,
            "draft_resume_path": self.draft_resume_path,
            "created_at": self.created_at,
            "source": self.source,
            "readiness_score": self.readiness_score,
            "keyword_gaps": self.keyword_gaps,
            "matched_keywords": self.matched_keywords,
            "suggestions": [s.to_dict() for s in self.suggestions],
        }

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> OptimizeBundle:
        return cls(
            job_id=data.get("job_id"),
            draft_resume_path=str(data.get("draft_resume_path") or ""),
            created_at=str(data.get("created_at") or ""),
            source=str(data.get("source") or "rule"),
            readiness_score=float(data.get("readiness_score") or 0.0),
            keyword_gaps=list(data.get("keyword_gaps") or []),
            matched_keywords=list(data.get("matched_keywords") or []),
            suggestions=[OptimizeSuggestion.from_dict(s) for s in (data.get("suggestions") or [])],
        )


def suggestions_json_path(draft_resume_path: Path | str) -> Path:
    path = Path(draft_resume_path)
    return path.with_name(f"{path.stem}_optimize.json")


def load_optimize_bundle(draft_resume_path: Path | str) -> OptimizeBundle | None:
    path = suggestions_json_path(draft_resume_path)
    if not path.exists():
        return None
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return OptimizeBundle.from_dict(data)
    except Exception as exc:
        logger.warning("Could not load optimize bundle %s: %s", path, exc)
        return None


def save_optimize_bundle(bundle: OptimizeBundle, draft_resume_path: Path | str | None = None) -> Path:
    draft = Path(draft_resume_path or bundle.draft_resume_path)
    out = suggestions_json_path(draft)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(bundle.to_dict(), indent=2), encoding="utf-8")
    return out


def _extract_bullets(resume_text: str) -> list[str]:
    bullets: list[str] = []
    for line in (resume_text or "").splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        match = _BULLET_RE.match(stripped)
        if match:
            bullets.append(match.group(1).strip())
        elif stripped.startswith(("•", "-", "*")) and len(stripped) > 8:
            bullets.append(stripped.lstrip("•-* ").strip())
    return bullets


def _rule_based_suggestions(
    job: ParsedJob,
    match: MatchExplanation,
    resume_text: str,
    gaps: list[str],
) -> list[OptimizeSuggestion]:
    suggestions: list[OptimizeSuggestion] = []
    bullets = _extract_bullets(resume_text)
    resume_norm = normalize_text(resume_text)
    gaps = filter_ats_keywords(gaps)

    # Keyword inserts: only for gaps already present somewhere in resume (truthful weave)
    weaveable = [g for g in gaps if _skill_in_text(g, resume_norm)][:5]
    for skill in weaveable:
        target = next((b for b in bullets if not _skill_in_text(skill, b)), bullets[0] if bullets else "")
        if not target:
            continue
        suggested = f"{target.rstrip('.')} using {skill}."
        suggestions.append(
            OptimizeSuggestion(
                id=str(uuid.uuid4()),
                kind="keyword_insert",
                status="pending",
                original_text=target,
                suggested_text=suggested,
                rationale=f"Emphasize existing keyword '{skill}' already present in your resume.",
                target_section="experience",
                keyword=skill,
            )
        )

    # Gaps not in resume: advisory only (do not invent as facts)
    advisory = [g for g in gaps if not _skill_in_text(g, resume_norm)][:5]
    for skill in advisory:
        suggestions.append(
            OptimizeSuggestion(
                id=str(uuid.uuid4()),
                kind="keyword_insert",
                status="pending",
                original_text="",
                suggested_text=f"[CONFIRM BEFORE ADDING] Add '{skill}' only if you have real experience with it.",
                rationale=f"Recruiter keyword '{skill}' appears in the job description but not in your resume.",
                target_section="skills",
                keyword=skill,
            )
        )

    # Bullet rewrites: achievement-oriented phrasing for first few bullets
    for bullet in bullets[:4]:
        if len(bullet) < 20:
            continue
        matched_bits = [s for s in match.matched_skills if _skill_in_text(s, bullet)][:2]
        suffix = f" aligned with {', '.join(matched_bits)}" if matched_bits else f" supporting {job.title}"
        suggested = bullet
        if not re.search(r"\b(led|built|improved|reduced|increased|delivered|designed)\b", bullet, re.I):
            suggested = f"Delivered {bullet[0].lower() + bullet[1:] if bullet else bullet}"
        if matched_bits and matched_bits[0].lower() not in suggested.lower():
            suggested = f"{suggested.rstrip('.')}{suffix}."
        if suggested.strip() == bullet.strip():
            continue
        suggestions.append(
            OptimizeSuggestion(
                id=str(uuid.uuid4()),
                kind="bullet_rewrite",
                status="pending",
                original_text=bullet,
                suggested_text=suggested,
                rationale="Achievement-oriented rewrite using role-aligned language (facts preserved).",
                target_section="experience",
            )
        )

    # Summary rewrite from opening lines
    opening = next((ln.strip() for ln in resume_text.splitlines() if len(ln.strip()) > 40), "")
    skills_phrase = ", ".join(match.matched_skills[:5]) or "core technical skills"
    summary = (
        f"{job.title}-focused professional with proven experience in {skills_phrase}. "
        f"Ready to contribute at {job.company}."
    )[:400]
    suggestions.append(
        OptimizeSuggestion(
            id=str(uuid.uuid4()),
            kind="summary_rewrite",
            status="pending",
            original_text=opening[:400],
            suggested_text=summary,
            rationale="ATS-oriented summary using matched skills already evidenced in your profile/resume.",
            target_section="summary",
        )
    )
    return suggestions


def _try_ollama_suggestions(
    job: ParsedJob,
    match: MatchExplanation,
    resume_text: str,
    settings: Settings,
) -> list[OptimizeSuggestion]:
    prompt = (
        "You are a truthful resume optimizer. Do NOT invent employers, dates, titles, or skills.\n"
        "Return ONLY a JSON array of objects with keys: kind, original_text, suggested_text, rationale, keyword.\n"
        "kind must be one of: keyword_insert, bullet_rewrite, summary_rewrite.\n"
        "Only rewrite or emphasize content that already exists in the resume.\n"
        f"Job title: {job.title}\nCompany: {job.company}\n"
        f"Matched skills: {', '.join(match.matched_skills)}\n"
        f"Missing skills: {', '.join(match.missing_skills)}\n"
        f"Job description (excerpt):\n{(job.description or '')[:2500]}\n\n"
        f"Resume text (excerpt):\n{(resume_text or '')[:3500]}\n"
    )
    raw = generate_ollama_completion(prompt, settings)
    if not raw:
        return []
    try:
        start = raw.find("[")
        end = raw.rfind("]")
        if start < 0 or end <= start:
            return []
        items = json.loads(raw[start : end + 1])
        out: list[OptimizeSuggestion] = []
        for item in items[:12]:
            if not isinstance(item, dict):
                continue
            kind = str(item.get("kind") or "bullet_rewrite")
            keyword = str(item.get("keyword") or "").strip()
            if kind == "keyword_insert" and keyword and not is_actionable_ats_keyword(keyword):
                continue
            out.append(
                OptimizeSuggestion(
                    id=str(uuid.uuid4()),
                    kind=kind,
                    status="pending",
                    original_text=str(item.get("original_text") or ""),
                    suggested_text=str(item.get("suggested_text") or ""),
                    rationale=str(item.get("rationale") or "AI suggestion — review for accuracy."),
                    keyword=keyword,
                    target_section=str(item.get("target_section") or ""),
                )
            )
        return out
    except Exception as exc:
        logger.debug("Ollama optimize parse failed: %s", exc)
        return []


def generate_optimize_suggestions(
    job: ParsedJob,
    match: MatchExplanation,
    master_resume_path: Path | str,
    draft_resume_path: Path | str,
    *,
    job_id: int | None = None,
    settings: Settings | None = None,
) -> OptimizeBundle:
    """Generate and persist optimize suggestions for a draft (Ollama + rule fallback)."""
    st = settings or get_settings()
    resume_text = extract_resume_text(master_resume_path)
    matched, gaps, readiness = compute_keyword_gap(job, resume_text, match)

    source = "rule"
    suggestions = _rule_based_suggestions(job, match, resume_text, gaps)
    ollama_items = _try_ollama_suggestions(job, match, resume_text, st)
    if ollama_items:
        # Prefer Ollama rewrites but keep advisory keyword gaps from rules
        advisory = [s for s in suggestions if s.suggested_text.startswith("[CONFIRM BEFORE ADDING]")]
        suggestions = ollama_items + advisory
        source = "hybrid" if suggestions else "ollama"

    # Final guard: never surface stop-word / noise keyword inserts in the UI.
    cleaned_suggestions: list[OptimizeSuggestion] = []
    for suggestion in suggestions:
        if suggestion.kind == "keyword_insert" and suggestion.keyword:
            if not is_actionable_ats_keyword(suggestion.keyword):
                continue
        cleaned_suggestions.append(suggestion)
    suggestions = cleaned_suggestions
    gaps = filter_ats_keywords(gaps)
    matched = filter_ats_keywords(matched)

    bundle = OptimizeBundle(
        job_id=job_id,
        draft_resume_path=str(draft_resume_path),
        created_at=datetime.now(timezone.utc).isoformat(),
        source=source,
        readiness_score=readiness,
        keyword_gaps=gaps,
        matched_keywords=matched,
        suggestions=suggestions,
    )
    save_optimize_bundle(bundle, draft_resume_path)
    logger.info(
        "Generated %d optimize suggestions for draft %s (source=%s, readiness=%.1f)",
        len(suggestions),
        draft_resume_path,
        source,
        readiness,
    )
    return bundle


def update_suggestion(
    draft_resume_path: Path | str,
    suggestion_id: str,
    action: str,
    edited_text: str | None = None,
) -> OptimizeBundle:
    """Accept, reject, or edit a single suggestion."""
    bundle = load_optimize_bundle(draft_resume_path)
    if not bundle:
        raise LookupError("No optimize suggestions found for this draft. Run AI Optimize first.")

    action_norm = action.strip().lower()
    found = False
    for suggestion in bundle.suggestions:
        if suggestion.id != suggestion_id:
            continue
        found = True
        if action_norm == "accept":
            suggestion.status = "accepted"
        elif action_norm == "reject":
            suggestion.status = "rejected"
        elif action_norm == "edit":
            if edited_text is None or not str(edited_text).strip():
                raise ValueError("edited_text is required for edit action")
            suggestion.suggested_text = str(edited_text).strip()
            suggestion.status = "edited"
        else:
            raise ValueError("action must be accept, reject, or edit")
        break

    if not found:
        raise LookupError(f"Suggestion id not found: {suggestion_id}")

    save_optimize_bundle(bundle, draft_resume_path)
    return bundle


def apply_accepted_suggestions_to_draft(
    draft_resume_path: Path | str,
    *,
    master_resume_path: Path | str | None = None,
) -> Path:
    """
    Apply accepted/edited suggestions into the draft DOCX.
    Keyword confirmations that start with [CONFIRM BEFORE ADDING] are skipped.
    """
    draft = Path(draft_resume_path)
    if not draft.exists():
        raise FileNotFoundError(f"Draft resume not found: {draft}")

    bundle = load_optimize_bundle(draft)
    if not bundle:
        raise LookupError("No optimize suggestions to apply.")

    accepted = [
        s
        for s in bundle.suggestions
        if s.status in {"accepted", "edited"}
        and s.suggested_text
        and not s.suggested_text.startswith("[CONFIRM BEFORE ADDING]")
    ]
    if not accepted:
        raise ValueError("No accepted suggestions to apply. Accept or edit at least one first.")

    source = Path(master_resume_path) if master_resume_path else draft
    if not source.exists():
        source = draft
    doc = Document(str(source if source.suffix.lower() == ".docx" else draft))

    # Replace paragraph text when original matches; append an Applied Optimizations section
    replacements = {
        s.original_text.strip(): s.suggested_text.strip()
        for s in accepted
        if s.original_text.strip() and s.kind in {"bullet_rewrite", "keyword_insert", "summary_rewrite"}
    }
    if replacements:
        for para in doc.paragraphs:
            text = (para.text or "").strip()
            for original, suggested in replacements.items():
                if original and original in text:
                    para.text = text.replace(original, suggested)

    doc.add_page_break()
    heading = doc.add_paragraph()
    run = heading.add_run("AI Optimize — Applied Suggestions")
    run.bold = True
    for suggestion in accepted:
        doc.add_paragraph(
            f"[{suggestion.kind}/{suggestion.status}] {suggestion.suggested_text}",
            style=None,
        )
        if suggestion.rationale:
            doc.add_paragraph(f"  Reason: {suggestion.rationale}")

    doc.save(str(draft))
    logger.info("Applied %d accepted optimize suggestions to %s", len(accepted), draft)
    return draft

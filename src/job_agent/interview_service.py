"""Service for tracking interviews and generating local truthful interview preparation documents."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from docx import Document
from sqlalchemy import select
from sqlalchemy.orm import Session

from job_agent.document_exporter import application_folder
from job_agent.database import InterviewRecord, JobRecord
from job_agent.logging_config import get_logger
from job_agent.models import CandidateProfile
from job_agent.resume_parser import extract_resume_text

logger = get_logger(__name__)


def add_interview(
    session: Session,
    job_id: int,
    interview_date: datetime,
    interview_type: str = "Screening",
    participants: str | None = None,
    notes: str | None = None,
    prep_tasks: str | None = None,
) -> InterviewRecord:
    rec = InterviewRecord(
        job_id=job_id,
        interview_date=interview_date,
        interview_type=interview_type,
        participants=participants,
        notes=notes,
        prep_tasks=prep_tasks,
    )
    session.add(rec)

    # Update job status to Interviewing if not already
    job = session.get(JobRecord, job_id)
    if job:
        job.status = "Interviewing"

    session.commit()
    session.refresh(rec)
    logger.info("Recorded interview for Job #%d on %s", job_id, interview_date)
    return rec


def list_interviews(session: Session, job_id: int | None = None) -> list[InterviewRecord]:
    stmt = select(InterviewRecord).order_by(InterviewRecord.interview_date.asc())
    if job_id is not None:
        stmt = stmt.where(InterviewRecord.job_id == job_id)
    return list(session.scalars(stmt))


def prepare_interview_doc(
    session: Session,
    job_id: int,
    profile: CandidateProfile,
    output_base_dir: Path,
) -> Path:
    """
    Generate a local, truthful interview preparation DOCX document using only the job description,
    candidate profile, master resume text, and user-entered notes.
    """
    job = session.get(JobRecord, job_id)
    if not job:
        raise ValueError(f"Job id {job_id} not found.")

    folder = application_folder(output_base_dir, job.company, job.title)
    doc_path = folder / f"{job.company}_{job.title}_Interview_Prep.docx".replace(" ", "_")

    doc = Document()
    doc.add_heading(f"Interview Preparation Sheet: {job.title} at {job.company}", level=1)

    doc.add_heading("1. Role Overview & Key Details", level=2)
    doc.add_paragraph(f"Company: {job.company}")
    doc.add_paragraph(f"Title: {job.title}")
    doc.add_paragraph(f"Location: {job.location or 'Not specified'}")
    doc.add_paragraph(f"Source URL: {job.job_url or 'N/A'}")

    doc.add_heading("2. Target Qualifications & Skills to Emphasize", level=2)
    matched_skills = [s.strip() for s in (job.matched_skills or "").split(",") if s.strip()]
    if matched_skills:
        doc.add_paragraph("Core Strengths & Matched Skills:\n - " + "\n - ".join(matched_skills))

    missing_skills = [s.strip() for s in (job.missing_skills or "").split(",") if s.strip()]
    if missing_skills:
        doc.add_paragraph("Potential Gap Areas / Follow-up Points to Address:\n - " + "\n - ".join(missing_skills))

    doc.add_heading("3. Candidate Resume & Background Reference", level=2)
    resume_path = profile.get_master_resume_path(job.title)
    resume_text = extract_resume_text(resume_path)
    if resume_text:
        snippet = resume_text[:1200] + ("..." if len(resume_text) > 1200 else "")
        doc.add_paragraph(f"Master Resume Snippet:\n{snippet}")
    else:
        doc.add_paragraph(f"Years of Experience: {profile.years_experience}")
        doc.add_paragraph(f"Target Roles: {', '.join(profile.target_titles)}")

    # Fetch recorded interviews & notes
    interviews = list_interviews(session, job_id)
    if interviews:
        doc.add_heading("4. Recorded Interview Schedule & Notes", level=2)
        for iv in interviews:
            doc.add_paragraph(f"• {iv.interview_type} on {iv.interview_date.strftime('%Y-%m-%d %H:%M')}")
            if iv.participants:
                doc.add_paragraph(f"  Participants: {iv.participants}")
            if iv.notes:
                doc.add_paragraph(f"  Notes: {iv.notes}")
            if iv.prep_tasks:
                doc.add_paragraph(f"  Tasks: {iv.prep_tasks}")

    if job.user_notes or job.notes:
        doc.add_heading("5. User Notes & Custom Instructions", level=2)
        doc.add_paragraph(job.user_notes or job.notes or "")

    doc.save(str(doc_path))
    logger.info("Saved interview preparation doc to: %s", doc_path)
    return doc_path
